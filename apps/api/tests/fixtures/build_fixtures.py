"""Generate the fixture contracts.

    uv run python tests/fixtures/build_fixtures.py

Five real-shaped agreements with the numbering conventions the chunker must
respect, one DOCX with heading styles, and one image-only PDF standing in for
a scan. Generated rather than downloaded so the suite is offline and the
expected structure is known exactly. The generated files are committed;
rerun this only when the templates change.
"""

# ruff: noqa: E501 — contract templates are long lines by nature
from __future__ import annotations

import textwrap
from pathlib import Path

import pymupdf
from docx import Document as DocxDocument

HERE = Path(__file__).parent

LOREM = (
    "The parties acknowledge that this provision has been negotiated at arm's length and "
    "reflects the commercial allocation of risk agreed between them. Nothing in this clause "
    "shall be construed to limit any liability which cannot be limited by applicable law. "
    "Each party shall bear its own costs in connection with the matters described herein."
)


def clause(number: str, title: str, body: str, subs: list[tuple[str, str]] | None = None) -> str:
    out = f"{number} {title}\n{body}\n"
    for label, text in subs or []:
        out += f"{label} {text}\n"
    return out


MSA = f"""MASTER SERVICES AGREEMENT

This Master Services Agreement (the "Agreement") is entered into as of 1 March 2025 between Aurora Holdings Limited, a company incorporated in England and Wales ("Customer"), and Boreal Software GmbH, a company incorporated in Germany ("Supplier").

ARTICLE I DEFINITIONS AND INTERPRETATION

1. Definitions
In this Agreement the following terms have the following meanings.
(a) "Affiliate" means any entity that controls, is controlled by, or is under common control with a party.
(b) "Confidential Information" means all information disclosed by one party to the other that is marked confidential or would reasonably be understood to be confidential.
(c) "Services" means the services described in a Statement of Work.

2. Interpretation
Clause headings are for convenience only and do not affect interpretation. Words in the singular include the plural and vice versa.

ARTICLE II SERVICES

3. Provision of Services
3.1 Supplier shall provide the Services in accordance with each Statement of Work and with reasonable skill and care.
3.2 Supplier shall comply with all applicable laws and with Customer's reasonable security policies notified to Supplier in writing.
3.3 Time shall not be of the essence in respect of any delivery date unless expressly stated in a Statement of Work.

4. Change Control
4.1 Either party may request a change to the Services by written notice.
4.2 No change shall be effective until agreed in writing by both parties.
{LOREM}

ARTICLE III COMMERCIAL TERMS

5. Fees and Payment
5.1 Customer shall pay the Fees set out in the applicable Statement of Work within thirty (30) days of receipt of a valid invoice.
5.2 All sums are exclusive of VAT, which shall be payable in addition at the applicable rate.
5.3 Supplier may charge interest on late payment at four per cent (4%) per annum above the Bank of England base rate.

6. Term and Termination
6.1 This Agreement commences on the Effective Date and continues for an initial term of three (3) years (the "Initial Term").
6.2 Following the Initial Term this Agreement shall renew automatically for successive periods of twelve (12) months unless either party gives not less than ninety (90) days' written notice of non-renewal.
6.3 Either party may terminate this Agreement for convenience on six (6) months' written notice.
6.4 Either party may terminate this Agreement immediately on written notice if the other party commits a material breach which is irremediable or which is not remedied within thirty (30) days of notice.

7. Confidentiality
7.1 Each party shall keep the other party's Confidential Information confidential and shall not disclose it except as permitted by this clause.
7.2 The obligations in this clause 7 survive termination for a period of five (5) years.
{LOREM}

ARTICLE IV RISK

8. Indemnities
8.1 Customer shall indemnify Supplier against all losses arising from Customer's breach of clause 3.2.
8.2 Supplier shall indemnify Customer against all losses arising from any claim that the Services infringe a third party's intellectual property rights.
8.3 The indemnified party shall notify the indemnifying party promptly of any claim and shall not settle any claim without the indemnifying party's prior written consent.

9. Limitation of Liability
9.1 Nothing in this Agreement limits liability for death or personal injury caused by negligence, for fraud, or for any liability which cannot be limited by law.
9.2 Subject to clause 9.1, each party's total aggregate liability under or in connection with this Agreement in any Contract Year shall not exceed an amount equal to one hundred and twenty-five per cent (125%) of the Fees paid or payable in that Contract Year.
9.3 Subject to clause 9.1, neither party shall be liable for any indirect or consequential loss, loss of profit, or loss of business.
{LOREM}
{LOREM}

ARTICLE V GENERAL

10. Assignment
Neither party may assign this Agreement without the prior written consent of the other, such consent not to be unreasonably withheld, save that Customer may assign to an Affiliate on notice.

11. Governing Law and Jurisdiction
This Agreement and any dispute arising out of it shall be governed by the laws of England and Wales, and the courts of England and Wales shall have exclusive jurisdiction.

IN WITNESS WHEREOF the parties have executed this Agreement.
"""

NDA = f"""MUTUAL NON-DISCLOSURE AGREEMENT

This Agreement is made on 12 June 2025 between Kestrel Analytics Inc. and Harrow & Vane LLP (each a "Party").

1. Purpose
The Parties wish to exchange Confidential Information for the purpose of evaluating a potential commercial relationship (the "Purpose").

2. Confidential Information
2.1 "Confidential Information" means any information disclosed by a Party in connection with the Purpose.
2.2 Confidential Information excludes information which:
(a) is or becomes public other than through breach of this Agreement;
(b) was lawfully in the receiving Party's possession before disclosure;
(c) is independently developed without reference to the disclosing Party's information; or
(d) is lawfully obtained from a third party free of any duty of confidence.

3. Obligations
3.1 The receiving Party shall use Confidential Information only for the Purpose.
3.2 The receiving Party shall not disclose Confidential Information to any person other than its Representatives who need to know it for the Purpose and who are bound by equivalent obligations.
{LOREM}

4. Term
This Agreement continues for two (2) years from the date above. The obligations of confidentiality survive for three (3) years after expiry.

5. Governing Law
This Agreement is governed by the laws of the State of New York. The Parties submit to the exclusive jurisdiction of the courts located in New York County.

6. No Licence
Nothing in this Agreement grants any licence or right in any Confidential Information other than as expressly set out herein.
"""

SAAS = f"""SOFTWARE AS A SERVICE SUBSCRIPTION AGREEMENT

Between Lumen Cloud Pty Ltd ("Provider") and Tessellate Retail Limited ("Subscriber"), dated 3 September 2025.

SECTION 1 SUBSCRIPTION
1.1 Provider grants Subscriber a non-exclusive, non-transferable right to access the Platform during the Subscription Term.
1.2 Subscriber shall not (a) resell the Platform, (b) reverse engineer any part of it, or (c) use it to build a competing product.

SECTION 2 DATA PROTECTION
2.1 Each party shall comply with Data Protection Laws.
2.2 Provider processes Subscriber Personal Data as processor on Subscriber's documented instructions, as set out in the Data Processing Addendum.
2.3 Provider shall not transfer Subscriber Personal Data outside the United Kingdom or the European Economic Area without Subscriber's prior written consent and appropriate safeguards.
{LOREM}

SECTION 3 EXCLUSIVITY
3.1 During the Subscription Term Subscriber shall not procure a substantially similar platform from any third party for the Territory.
3.2 The restriction in clause 3.1 falls away if Provider fails to meet the Service Levels in any three (3) consecutive months.

SECTION 4 INTELLECTUAL PROPERTY
4.1 Provider retains all intellectual property rights in the Platform.
4.2 Subscriber retains all intellectual property rights in Subscriber Data and grants Provider a licence to use it solely to provide the Platform.
4.3 Any improvements to the Platform arising from Subscriber feedback shall vest in Provider.

SECTION 5 CHANGE OF CONTROL
Provider may terminate this Agreement on thirty (30) days' notice if Subscriber undergoes a Change of Control in favour of a competitor of Provider.

SECTION 6 FEES
6.1 Fees are payable annually in advance.
6.2 Provider may increase the Fees on each anniversary by no more than the greater of five per cent (5%) and the increase in CPI.
{LOREM}

SECTION 7 GOVERNING LAW
This Agreement is governed by the laws of New South Wales, Australia.
"""

EMPLOYMENT = f"""EXECUTIVE EMPLOYMENT AGREEMENT

This Agreement is made on 15 January 2025 between Meridian Logistics plc (the "Company") and Dr Imogen Reyes (the "Executive").

1. Appointment
1.1 The Company appoints the Executive as Chief Operating Officer with effect from 1 February 2025.
1.2 The Executive's employment is subject to a probationary period of six (6) months.

2. Duties
2.1 The Executive shall devote the whole of her working time to the business of the Company.
2.2 The Executive shall not without prior written consent be engaged in any other business.

3. Remuneration
3.1 The Company shall pay the Executive a salary of £245,000 per annum, payable monthly in arrears.
3.2 The Executive is eligible for an annual bonus of up to fifty per cent (50%) of salary at the discretion of the Remuneration Committee.
{LOREM}

4. Termination
4.1 Either party may terminate the employment on not less than nine (9) months' written notice.
4.2 The Company may terminate summarily for gross misconduct.
4.3 The Company may at its discretion make a payment in lieu of notice equal to basic salary for the notice period.

5. Restrictive Covenants
5.1 For twelve (12) months after Termination the Executive shall not be engaged in any Competing Business within the United Kingdom.
5.2 For twelve (12) months after Termination the Executive shall not solicit any Restricted Customer.
5.3 The Executive acknowledges that the restrictions in this clause 5 are reasonable and necessary to protect the Company's legitimate business interests.
{LOREM}

6. Intellectual Property
All Intellectual Property created by the Executive in the course of employment shall vest in the Company.

7. Governing Law
This Agreement is governed by the laws of England and Wales.
"""

LEASE = f"""COMMERCIAL LEASE

This Lease is made on 30 April 2025 between Threadneedle Estates Limited (the "Landlord") and Coppice Coffee Roasters Limited (the "Tenant").

PART 1 PARTICULARS
1. Premises
The ground floor and basement of 14 Salter Row, Leeds, shown edged red on the Plan.
2. Term
Ten (10) years from and including 1 May 2025.
3. Rent
£48,000 per annum, payable quarterly in advance on the usual quarter days.

PART 2 TENANT COVENANTS
4. Rent
4.1 The Tenant shall pay the Rent without deduction or set-off.
4.2 Interest at four per cent (4%) above base rate accrues on any Rent unpaid for more than fourteen (14) days.
5. Repair
The Tenant shall keep the Premises in good and substantial repair and condition.
6. Alienation
6.1 The Tenant shall not assign the whole of the Premises without the Landlord's consent, not to be unreasonably withheld.
6.2 The Tenant shall not assign or underlet part only of the Premises.
6.3 The Tenant shall not charge the Premises.
{LOREM}

PART 3 BREAK AND RENEWAL
7. Tenant Break
7.1 The Tenant may terminate this Lease on the fifth anniversary of the Term Commencement Date by giving not less than six (6) months' written notice.
7.2 The break is conditional on the Rent being paid up to date and vacant possession being given.

PART 4 GENERAL
8. Dispute Resolution
Any dispute as to the Rent payable on review shall be referred to an independent surveyor acting as expert. All other disputes shall be determined by the courts of England and Wales.
9. Governing Law
This Lease is governed by the laws of England and Wales.
"""

TEMPLATES = {
    "msa.pdf": MSA,
    "nda.pdf": NDA,
    "saas_subscription.pdf": SAAS,
    "employment.pdf": EMPLOYMENT,
    "lease.pdf": LEASE,
}


def write_pdf(path: Path, body: str) -> None:
    doc = pymupdf.open()
    width, height = 595, 842  # A4 points
    margin = 56
    line_height = 14
    max_chars = 92

    lines: list[str] = []
    for paragraph in body.strip().split("\n"):
        if not paragraph.strip():
            lines.append("")
            continue
        lines.extend(textwrap.wrap(paragraph, max_chars) or [""])
        lines.append("")

    page = doc.new_page(width=width, height=height)
    y = margin
    for line in lines:
        if y > height - margin:
            page = doc.new_page(width=width, height=height)
            y = margin
        if line:
            page.insert_text((margin, y), line, fontsize=10, fontname="helv")
        y += line_height
    doc.save(path)
    doc.close()


def write_scanned_pdf(path: Path, source: Path) -> None:
    """Rasterise the first two pages of ``source`` so the PDF has no text layer."""
    src = pymupdf.open(source)
    out = pymupdf.open()
    for index in range(min(2, src.page_count)):
        page = src[index]
        pix = page.get_pixmap(dpi=100, colorspace=pymupdf.csGRAY)
        new_page = out.new_page(width=page.rect.width, height=page.rect.height)
        # JPEG keeps the fixture small enough to commit.
        new_page.insert_image(new_page.rect, stream=pix.tobytes("jpeg", jpg_quality=55))
    out.save(path, deflate=True, garbage=3)
    out.close()
    src.close()


def write_docx(path: Path, body: str) -> None:
    doc = DocxDocument()
    for paragraph in body.strip().split("\n"):
        stripped = paragraph.strip()
        if not stripped:
            continue
        if stripped.isupper() and len(stripped) < 80:
            doc.add_heading(stripped, level=1)
        elif stripped[:2].rstrip(".").isdigit() and "." not in stripped[:3]:
            doc.add_heading(stripped, level=2)
        else:
            doc.add_paragraph(stripped)
    doc.save(str(path))


def main() -> None:
    for name, body in TEMPLATES.items():
        write_pdf(HERE / name, body)
    write_scanned_pdf(HERE / "scanned_nda.pdf", HERE / "nda.pdf")
    write_docx(HERE / "msa.docx", MSA)
    print("fixtures written to", HERE)


if __name__ == "__main__":
    main()
