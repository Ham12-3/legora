"""Findings as a DOCX issues list, the shape lawyers circulate."""

import io
from datetime import UTC, datetime

from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from app.models.playbook import Finding, MatchedPosition, Severity

_SEVERITY_ORDER = {Severity.HIGH: 0, Severity.MEDIUM: 1, Severity.LOW: 2, Severity.NONE: 3}
_POSITION_LABEL = {
    MatchedPosition.PREFERRED: "Preferred",
    MatchedPosition.FALLBACK: "Fallback",
    MatchedPosition.UNACCEPTABLE: "Unacceptable",
    MatchedPosition.NOT_ADDRESSED: "Not addressed",
}


def issues_list_docx(
    *, playbook_name: str, document_name: str, findings: list[Finding], model: str | None
) -> bytes:
    doc = DocxDocument()
    doc.add_heading("Issues List", level=0)
    meta = doc.add_paragraph()
    meta.add_run(f"Document: {document_name}\n").bold = True
    meta.add_run(f"Playbook: {playbook_name}\n")
    meta.add_run(f"Generated: {datetime.now(UTC).strftime('%d %B %Y, %H:%M UTC')}\n")
    meta.add_run(
        "AI-assisted review. Every quotation below was matched to the source text; "
        "items marked UNVERIFIED could not be, and must be checked by hand. "
        "This is not legal advice."
    ).italic = True

    ordered = sorted(findings, key=lambda f: (_SEVERITY_ORDER[f.severity], f.ordinal))
    deviations = [f for f in ordered if f.matched_position is not MatchedPosition.PREFERRED]
    doc.add_paragraph(
        f"{len(findings)} topics reviewed; {len(deviations)} deviate from the preferred position."
    )

    table = doc.add_table(rows=1, cols=6)
    table.style = "Light Grid Accent 1"
    for cell, title in zip(
        table.rows[0].cells,
        ["#", "Topic", "Clause", "Position", "Severity", "Issue and proposed language"],
        strict=True,
    ):
        cell.text = title
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True

    for n, f in enumerate(ordered, start=1):
        row = table.add_row().cells
        row[0].text = str(n)
        row[1].text = f.topic
        row[2].text = f.clause_reference or "—"
        row[3].text = _POSITION_LABEL[f.matched_position]
        row[4].text = f.severity.value.upper()
        body = row[5].paragraphs[0]
        if f.quoted_text:
            q = body.add_run(f"“{f.quoted_text}”")
            q.italic = True
            body.add_run(f" (p. {f.page})\n")
        elif f.matched_position is not MatchedPosition.NOT_ADDRESSED:
            body.add_run("UNVERIFIED — no quote could be matched to the document.\n").bold = True
        body.add_run(f.rationale)
        if f.suggested_language:
            p2 = row[5].add_paragraph()
            p2.add_run("Proposed language: ").bold = True
            p2.add_run(f.suggested_language)

    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.size = Pt(9)

    footer = doc.add_paragraph()
    footer.add_run(f"Model: {model or 'n/a'}").font.size = Pt(8)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
