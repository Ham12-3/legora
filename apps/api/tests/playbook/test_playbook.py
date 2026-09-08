# ruff: noqa: E501 -- playbook positions are long sentences by nature
"""Playbooks against the database with the fake model."""

import io
import uuid
from pathlib import Path

import httpx
from docx import Document as DocxDocument

from app.db import SessionLocal
from app.ingestion import pipeline
from app.ingestion.embeddings import FakeEmbedder
from app.llm.fake import FakeLLMClient
from app.playbook import service
from tests.conftest import Actor, MakeActor, create_document, create_matter

FIXTURES = Path(__file__).parent.parent / "fixtures"

RULES = [
    {
        "topic": "Governing law",
        "preferred_position": "The agreement is governed by the laws of England and Wales.",
        "fallback_position": "The agreement is governed by the laws of Scotland or Ireland.",
        "unacceptable_position": "The agreement is governed by the laws of New York or Delaware.",
    },
    {
        "topic": "Liability cap",
        "preferred_position": (
            "Aggregate liability capped at one hundred per cent of the fees paid in the contract year."
        ),
        "fallback_position": "Aggregate liability capped at one hundred and twenty-five per cent of fees.",
        "unacceptable_position": "Unlimited liability or a cap above two hundred per cent of fees.",
    },
    {
        "topic": "Termination for convenience",
        "preferred_position": "Either party may terminate for convenience on three months written notice.",
        "fallback_position": "Either party may terminate for convenience on six months written notice.",
        "unacceptable_position": "No right to terminate for convenience.",
    },
    {
        "topic": "Source code escrow",
        "preferred_position": "Supplier deposits source code in escrow with release on insolvency.",
        "fallback_position": "Supplier commits to negotiate escrow in good faith.",
        "unacceptable_position": "No escrow arrangement.",
    },
]


def loader_for(name: str) -> pipeline.Loader:
    data = (FIXTURES / name).read_bytes()

    async def _load(_: str) -> bytes:
        return data

    return _load


async def ingest(document_id: uuid.UUID, fixture: str) -> None:
    async with SessionLocal() as session:
        await pipeline.run_all(
            session, document_id, loader=loader_for(fixture), embedder=FakeEmbedder(1536)
        )


async def setup(
    client: httpx.AsyncClient, actor: Actor, fixture: str = "msa.pdf"
) -> tuple[uuid.UUID, uuid.UUID]:
    matter_id = await create_matter(client, actor)
    doc_id = await create_document(client, actor, matter_id, filename=fixture)
    await ingest(doc_id, fixture)
    r = await client.post(
        "/playbooks",
        json={"name": "Supplier paper", "description": "Our standard positions", "rules": RULES},
        headers=actor.headers,
    )
    assert r.status_code == 201, r.text
    assert len(r.json()["rules"]) == len(RULES)
    return uuid.UUID(r.json()["playbook"]["id"]), doc_id


async def run(
    client: httpx.AsyncClient,
    actor: Actor,
    playbook_id: uuid.UUID,
    doc_id: uuid.UUID,
    fake: FakeLLMClient,
) -> dict[str, object]:
    r = await client.post(
        f"/playbooks/{playbook_id}/run", json={"document_id": str(doc_id)}, headers=actor.headers
    )
    assert r.status_code == 202, r.text
    run_id = uuid.UUID(r.json()["id"])
    async with SessionLocal() as session:
        await service.run_playbook(session, run_id, client=fake, embedder=None)
    detail = await client.get(f"/playbook-runs/{run_id}", headers=actor.headers)
    assert detail.status_code == 200, detail.text
    body: dict[str, object] = detail.json()
    return body


async def test_findings_cover_every_rule_with_verified_citations(
    client: httpx.AsyncClient, make_actor: MakeActor
) -> None:
    alice = await make_actor("alice")
    playbook_id, doc_id = await setup(client, alice)
    fake = FakeLLMClient()

    body = await run(client, alice, playbook_id, doc_id, fake)
    run_out = body["run"]
    assert isinstance(run_out, dict) and run_out["status"] == "done"
    assert run_out["model"] == "fake" and body["demo_mode"] is True
    findings = body["findings"]
    assert isinstance(findings, list) and len(findings) == len(RULES)
    by_topic = {f["topic"]: f for f in findings}

    addressed = [f for f in findings if f["matched_position"] != "not_addressed"]
    assert addressed, "the MSA addresses governing law, liability and termination"
    for f in addressed:
        assert f["verified"] is True, f
        assert f["quoted_text"] and f["page"] >= 1 and f["bboxes"], f
        assert f["match_kind"] in {"exact", "relocated"}
        assert f["rationale"]

    escrow = by_topic["Source code escrow"]
    assert escrow["matched_position"] == "not_addressed"
    assert escrow["verified"] is True and escrow["quoted_text"] is None
    assert escrow["severity"] == "medium"
    assert escrow["suggested_language"]

    # Deviations carry replacement language; a preferred match does not.
    for f in addressed:
        if f["matched_position"] == "preferred":
            assert f["suggested_language"] == "" and f["severity"] == "none"
        else:
            assert f["suggested_language"]


async def test_fabricated_quotes_leave_findings_unverified(
    client: httpx.AsyncClient, make_actor: MakeActor
) -> None:
    alice = await make_actor("alice")
    playbook_id, doc_id = await setup(client, alice)
    body = await run(client, alice, playbook_id, doc_id, FakeLLMClient(fabricate=True))
    findings = body["findings"]
    assert isinstance(findings, list)
    addressed = [f for f in findings if f["matched_position"] != "not_addressed"]
    assert addressed
    for f in addressed:
        assert f["verified"] is False
        assert f["quoted_text"] is None, "unverified quotes are never stored"


async def test_export_docx_lists_findings(client: httpx.AsyncClient, make_actor: MakeActor) -> None:
    alice = await make_actor("alice")
    playbook_id, doc_id = await setup(client, alice)
    body = await run(client, alice, playbook_id, doc_id, FakeLLMClient())
    run_out = body["run"]
    assert isinstance(run_out, dict)

    r = await client.get(f"/playbook-runs/{run_out['id']}/export", headers=alice.headers)
    assert r.status_code == 200
    assert r.content[:2] == b"PK"
    assert "issues-" in r.headers["content-disposition"]
    docx = DocxDocument(io.BytesIO(r.content))
    text = "\n".join(p.text for p in docx.paragraphs)
    assert "Issues List" in text and "msa.pdf" in text
    table = docx.tables[0]
    assert len(table.rows) == len(RULES) + 1
    cells = "\n".join(c.text for row in table.rows for c in row.cells)
    assert "Source code escrow" in cells and "Not addressed" in cells
    assert "Proposed language" in cells


async def test_run_lists_and_not_ready_document_fails_clearly(
    client: httpx.AsyncClient, make_actor: MakeActor
) -> None:
    alice = await make_actor("alice")
    matter_id = await create_matter(client, alice)
    pending = await create_document(client, alice, matter_id, filename="pending.pdf")
    r = await client.post(
        "/playbooks", json={"name": "P", "rules": RULES[:1]}, headers=alice.headers
    )
    playbook_id = r.json()["playbook"]["id"]
    r = await client.post(
        f"/playbooks/{playbook_id}/run", json={"document_id": str(pending)}, headers=alice.headers
    )
    run_id = uuid.UUID(r.json()["id"])
    async with SessionLocal() as session:
        await service.run_playbook(session, run_id, client=FakeLLMClient(), embedder=None)
    listed = await client.get(f"/playbook-runs?document_id={pending}", headers=alice.headers)
    assert [x["id"] for x in listed.json()] == [str(run_id)]
    assert listed.json()[0]["status"] == "failed"
    assert "not ready" in listed.json()[0]["error"]


async def test_rule_crud(client: httpx.AsyncClient, make_actor: MakeActor) -> None:
    alice = await make_actor("alice")
    r = await client.post("/playbooks", json={"name": "Empty"}, headers=alice.headers)
    playbook_id = r.json()["playbook"]["id"]
    r = await client.post(f"/playbooks/{playbook_id}/rules", json=RULES[0], headers=alice.headers)
    assert r.status_code == 201
    rule_id = r.json()["id"]
    r = await client.patch(
        f"/playbooks/{playbook_id}/rules/{rule_id}", json={"topic": "Law"}, headers=alice.headers
    )
    assert r.json()["topic"] == "Law"
    r = await client.post(
        f"/playbooks/{playbook_id}/run",
        json={"document_id": str(uuid.uuid4())},
        headers=alice.headers,
    )
    assert r.status_code == 404
    r = await client.delete(f"/playbooks/{playbook_id}/rules/{rule_id}", headers=alice.headers)
    assert r.status_code == 204
    detail = await client.get(f"/playbooks/{playbook_id}", headers=alice.headers)
    assert detail.json()["rules"] == []
    r = await client.post(
        f"/playbooks/{playbook_id}/run",
        json={"document_id": str(uuid.uuid4())},
        headers=alice.headers,
    )
    assert r.status_code in (404, 409)
