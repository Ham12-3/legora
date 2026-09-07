"""CSV and XLSX export of a grid.

Unverified and not-found answers are marked in the export itself. A lawyer
reading a spreadsheet must not be able to mistake an unverified answer for a
verified one, any more than in the UI.
"""

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models.review import Cell, CellStatus, ReviewColumn

UNVERIFIED_SUFFIX = " [UNVERIFIED]"
NOT_FOUND = "Not found"


def display_value(cell: Cell | None) -> str:
    if cell is None or cell.status is not CellStatus.DONE:
        return "" if cell is None else f"[{cell.status.value}]"
    if cell.not_found:
        return NOT_FOUND
    text = cell.value_text or ""
    return text if cell.verified else text + UNVERIFIED_SUFFIX


def build_table(
    rows: list[tuple[str, dict[str, Cell]]], columns: list[ReviewColumn]
) -> tuple[list[str], list[list[str]]]:
    """``rows`` is [(document filename, {column_id: cell})]."""
    headers = ["Document", *[c.name for c in columns]]
    body = [
        [filename, *[display_value(cells.get(str(c.id))) for c in columns]]
        for filename, cells in rows
    ]
    return headers, body


def to_csv(headers: list[str], body: list[list[str]]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(headers)
    writer.writerows(body)
    return buffer.getvalue().encode("utf-8-sig")


def to_xlsx(headers: list[str], body: list[list[str]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Review"
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    amber = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
    for row in body:
        sheet.append(row)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str) and cell.value.endswith(UNVERIFIED_SUFFIX):
                cell.fill = amber
    sheet.freeze_panes = "B2"
    for index, column_cells in enumerate(sheet.columns, start=1):
        width = max(len(str(c.value or "")) for c in column_cells)
        sheet.column_dimensions[get_column_letter(index)].width = min(max(12, width + 2), 60)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
